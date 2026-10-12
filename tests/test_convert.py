"""Conversión de Word, TXT, FB2 y Kindle a EPUB (textos inventados)."""

import base64
import os
import stat
import sys

import pytest
from helpers import PNG, make_epub
from lxml import etree

from traductor import convert as conv
from traductor.convert import ConvertError, convert
from traductor.epub import EpubBook


def texts(path):
    with EpubBook(path) as book:
        out = []
        for doc in book.documents:
            if doc.in_spine:
                out += [(s.tag, s.source) for s in book.segments(doc)]
        return out, book


def test_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.core_properties.title = "El faro del norte"
    doc.core_properties.author = "Autora Inventada"
    doc.add_paragraph("El faro del norte", style="Title")
    doc.add_heading("Capítulo uno", level=1)
    p = doc.add_paragraph("Mirela miró el mar ")
    p.add_run("gris").italic = True
    p.add_run(" toda la tarde.")
    doc.add_paragraph("Una cita del viejo farero.", style="Quote")
    doc.add_heading("Capítulo dos", level=1)
    doc.add_paragraph("Llegó la marea.")
    src = tmp_path / "libro.docx"
    doc.save(str(src))

    out, report = convert(src)
    assert out == tmp_path / "libro.epub" and report.title == "El faro del norte" and report.author == "Autora Inventada"
    assert report.chapters == 2 and report.language == "es"
    segs, _ = texts(out)
    assert ("h1", "Capítulo uno") in segs and ("h1", "Capítulo dos") in segs
    assert any("<g1>gris</g1>" in s for _t, s in segs)
    assert any(t == "p" and "Una cita del viejo farero." in s for t, s in segs)


def test_docx_footnotes_from_xml():
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    para = etree.fromstring(
        f'<w:p xmlns:w="{W}"><w:r><w:t>Texto con nota</w:t></w:r>'
        f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="2"/></w:r>'
        f'<w:r><w:rPr><w:b/></w:rPr><w:t xml:space="preserve"> y fin.</w:t></w:r></w:p>')
    notes = {}
    runs, refs = conv._docx_runs(para, {"2": [conv.Run("La nota.")]}, notes, [0])
    assert [(r.text, r.bold, r.sup) for r in runs] == [("Texto con nota", False, False), ("1", False, True), (" y fin.", True, False)]
    assert refs == [("1", "n1")] and notes == {"n1": [conv.Run("La nota.")]}


TXT = """EL FARO DEL NORTE


CHAPTER I

Mirela walked to the harbor
before dawn, as she did _every_ day.

The keeper waited for her.

* * *

She nodded and looked at the water.

CHAPTER II

The boats came back at noon.
"""


def test_txt(tmp_path):
    src = tmp_path / "faro.txt"
    src.write_text(TXT, encoding="cp1252")
    out, report = convert(src)
    segs, _ = texts(out)
    assert report.chapters == 3 and report.title == "faro" and report.language == "en"
    assert ("h1", "CHAPTER I") in segs and ("h1", "CHAPTER II") in segs
    assert ("p", "Mirela walked to the harbor before dawn, as she did <g1>every</g1> day.") in segs
    import zipfile

    with zipfile.ZipFile(out) as z:  # el separador no tiene letras: no se traduce, pero está en el libro
        assert any('class="centrado">* * *<' in z.read(n).decode() for n in z.namelist() if n.endswith(".xhtml"))



def test_txt_lines_without_blank_separators(tmp_path):
    src = tmp_path / "lineas.txt"
    src.write_text("Capítulo 1\nPrimera línea del libro.\nSegunda línea del libro.\nCapítulo 2\nOtra más.\n", encoding="utf-8")
    blocks, *_ = conv.txt_blocks(src)
    assert [(b.kind, b.runs[0].text) for b in blocks] == [
        ("h", "Capítulo 1"), ("p", "Primera línea del libro."), ("p", "Segunda línea del libro."),
        ("h", "Capítulo 2"), ("p", "Otra más.")]


FB2 = """<?xml version="1.0" encoding="utf-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
<description><title-info>
<author><first-name>Autora</first-name><last-name>Inventada</last-name></author>
<book-title>Redes de niebla</book-title><lang>es</lang>
<coverpage><image l:href="#portada.png"/></coverpage>
</title-info></description>
<body>
<section><title><p>Primera parte</p></title>
<p>Mirela remendaba las <emphasis>redes</emphasis> en el muelle<a l:href="#n1" type="note">[1]</a>.</p>
<epigraph><p>Una cita breve.</p></epigraph>
<section><title><p>La tormenta</p></title><p>Llovió <strong>tres</strong> días.</p></section>
</section>
<section><title><p>Segunda parte</p></title><p>Volvieron los barcos.</p></section>
</body>
<body name="notes"><section id="n1"><title><p>1</p></title><p>Nota inventada.</p></section></body>
<binary id="portada.png" content-type="image/png">""" + base64.b64encode(PNG).decode() + """</binary>
</FictionBook>"""


def test_fb2(tmp_path):
    src = tmp_path / "redes.fb2"
    src.write_text(FB2, encoding="utf-8")
    out, report = convert(src)
    assert (report.title, report.author, report.language) == ("Redes de niebla", "Autora Inventada", "es")
    assert report.chapters == 2 and report.notes == 1
    segs, _ = texts(out)
    assert ("h1", "Primera parte") in segs and ("h2", "La tormenta") in segs and ("h1", "Segunda parte") in segs
    assert any("<g1>redes</g1>" in s and "<x2/>" in s for _t, s in segs)  # cursiva y llamada a nota
    assert any("<g1>tres</g1>" in s for _t, s in segs)
    with EpubBook(out) as book:
        assert book.cover_item() is not None


def test_kindle_without_calibre(tmp_path, monkeypatch):
    src = tmp_path / "libro.azw3"
    src.write_bytes(b"BOOKMOBI")
    monkeypatch.setattr(conv, "find_calibre", lambda: None)
    with pytest.raises(ConvertError, match="Calibre"):
        convert(src)


@pytest.mark.skipif(sys.platform.startswith("win"), reason="script de shell")
def test_kindle_with_calibre(tmp_path, monkeypatch):
    epub = make_epub(tmp_path / "hecho.epub", [("One", "<p>Mirela walked.</p>")])
    fake = tmp_path / "ebook-convert"
    fake.write_text(f'#!/bin/sh\ncp "{epub}" "$2"\n')
    fake.chmod(fake.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("PATH", f"{tmp_path}{os.pathsep}{os.environ['PATH']}")
    src = tmp_path / "libro.mobi"
    src.write_bytes(b"BOOKMOBI")
    out, _report = convert(src)
    assert out == tmp_path / "libro.epub" and out.read_bytes() == epub.read_bytes()


def test_rejected_formats(tmp_path):
    for name, msg in (("x.odt", "no admitido"), ("x.epub", "ya es un EPUB")):
        (tmp_path / name).write_bytes(b"x")
        with pytest.raises(ConvertError, match=msg):
            convert(tmp_path / name)


def test_api_convert_book(tmp_path, monkeypatch):
    from test_app_api import wait

    from traductor.app.api import Api

    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    src = tmp_path / "faro.txt"
    src.write_text(TXT, encoding="utf-8")
    api = Api(mode="navegador")
    api.convert_book(str(src))
    st = wait(api)
    assert st["error"] is None and st["resultado"]["formato"] == "texto" and st["resultado"]["capitulos"] == 3
    with pytest.raises(ValueError, match="ya es un EPUB"):
        api.convert_book(str(make_epub(tmp_path / "a.epub", [("One", "<p>x</p>")])))
