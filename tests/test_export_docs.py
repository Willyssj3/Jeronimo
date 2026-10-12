"""Exportar a Word y PDF (libro y frases inventadas)."""

import pytest
from test_pipeline import project  # noqa: F401  (fixture)

from traductor import export_docs
from traductor.export_docs import Run, to_runs


def test_runs_keep_formatting_and_note_calls():
    runs = to_runs("Hola <g1>mundo &amp; <g2>mar</g2></g1><x3/>.<x4/>fin",
                   {1: ("em", ""), 2: ("b", ""), 3: ("a", "12"), 4: ("br", "")})
    assert runs == [Run("Hola "), Run("mundo & ", italic=True), Run("mar", italic=True, bold=True),
                    Run("12", sup=True), Run("."), Run("", br=True), Run("fin")]


@pytest.fixture
def translated(project):  # noqa: F811
    project.set_translation(0, 0, "El puerto")
    project.set_translation(0, 1, "Mirela caminó hasta Puertoarado antes del alba, como cada <g1>santo</g1> día.")
    project.set_translation(0, 2, "—Los barcos llegan tarde —dijo el farero<x1/>.")
    return project


def test_collect_uses_translation_or_original(translated):
    chapters = export_docs.collect(translated, with_source=True)
    assert [c.title for c in chapters] == ["The Harbor", "Second"]  # sin índices
    blocks = chapters[0].blocks
    assert blocks[0].level == 1 and blocks[0].runs == [Run("El puerto")]
    assert Run("santo", italic=True) in blocks[1].runs
    assert blocks[2].runs[-2:] == [Run("1", sup=True), Run(".")]
    assert not blocks[3].translated and export_docs._plain(blocks[3].runs).startswith("She nodded")
    assert export_docs._plain(blocks[1].src_runs).startswith("Mirela walked")
    translated.set_skipped(1)
    assert [c.title for c in export_docs.collect(translated)] == ["The Harbor"]


def test_docx(translated, tmp_path):
    from docx import Document

    out = export_docs.export_docx(translated, tmp_path / "libro.docx")
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts[0] == "A Test Book" and "El puerto" in texts
    para = next(p for p in doc.paragraphs if p.text.startswith("Mirela caminó"))
    assert any(r.italic and r.text == "santo" for r in para.runs)
    assert doc.paragraphs[0].style.name == "Title"
    assert any(p.style.name == "Heading 1" and p.text == "El puerto" for p in doc.paragraphs)

    out = export_docs.export_docx(translated, tmp_path / "bilingue.docx", bilingual=True)
    table = Document(str(out)).tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Original", "Traducción"]
    assert table.rows[2].cells[0].text.startswith("Mirela walked") and table.rows[2].cells[1].text.startswith("Mirela caminó")
    assert table.rows[4].cells[1].text == ""  # sin traducir: celda vacía


def test_pdf(translated, tmp_path):
    import pdfplumber

    out = export_docs.export_pdf(translated, tmp_path / "libro.pdf", page="a5")
    with pdfplumber.open(out) as pdf:
        assert len(pdf.pages) >= 3  # portadilla y un capítulo por página
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        width_mm = pdf.pages[0].width / 72 * 25.4
    assert "A Test Book" in text and "Puertoarado antes del alba" in text and "farero" in text
    assert abs(width_mm - 148) < 1

    translated.set_translation(0, 3, "東京の夜")
    with pytest.raises(ValueError, match="no tiene estos caracteres"):
        export_docs.export_pdf(translated, tmp_path / "x.pdf")


def test_api_export_document(translated, tmp_path, monkeypatch):
    from test_app_api import wait

    from traductor.app.api import Api

    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    api = Api(mode="navegador")
    api.open_project(str(translated.root))
    api.export_document("docx-bilingue")
    st = wait(api)
    assert st["error"] is None and st["resultado"]["ruta"].endswith(".bilingue-en-es.docx")
    api.export_document("pdf-a4")
    st = wait(api)
    assert st["error"] is None and st["resultado"]["ruta"].endswith(".es.a4.pdf")
    with pytest.raises(ValueError):
        api.export_document("odt")
