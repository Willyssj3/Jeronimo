"""Firma de Jerónimo al final de los libros exportados (EPUB, Word y PDF), con un libro inventado."""

import zipfile

import pytest
from lxml import etree
from test_pipeline import project  # noqa: F401  (fixture)

from traductor import credit, export_docs
from traductor.epub import EpubBook

OPF_NS = {"opf": "http://www.idpf.org/2007/opf"}


@pytest.fixture
def translated(project):  # noqa: F811
    project.set_translation(0, 0, "El puerto")
    project.set_translation(0, 1, "Mirela caminó hasta Puertoarado antes del alba.")
    return project


def test_texts_follow_the_book_language():
    assert credit.texts("es")[0] == "Traducido con Jerónimo"
    assert credit.texts("pt-BR")[0] == "Traduzido com o Jerónimo"
    assert credit.texts("ja")[0] == "Translated with Jerónimo"  # sin traducción propia: en inglés
    for epub3 in (True, False):
        root = etree.fromstring(credit.xhtml("es", epub3=epub3))  # XHTML bien formado
        assert credit.URL in etree.tostring(root, encoding="unicode")
    assert credit.logo_png()[:8] == b"\x89PNG\r\n\x1a\n"


def opf_of(epub):
    with zipfile.ZipFile(epub) as z:
        container = etree.fromstring(z.read("META-INF/container.xml"))
        path = container.find(".//{*}rootfile").get("full-path")
        return z.namelist(), etree.fromstring(z.read(path)), z


def test_epub_ends_with_the_signature_page(translated):
    out, _ = translated.export()
    names, opf, _z = opf_of(out)
    assert any(n.endswith("jeronimo-firma.xhtml") for n in names)
    assert any(n.endswith("jeronimo-logo.png") for n in names)
    spine = [ref.get("idref") for ref in opf.findall(".//opf:spine/opf:itemref", OPF_NS)]
    assert spine[-1] == "jeronimo-firma"
    with zipfile.ZipFile(out) as z:
        page = z.read(next(n for n in names if n.endswith("jeronimo-firma.xhtml"))).decode("utf-8")
    assert "Traducido con Jerónimo" in page and "github.com/Willyssj3/Jeronimo" in page
    # Un libro que ya tiene la firma no se firma dos veces.
    again = out.with_name("otra.epub")
    with EpubBook(out) as book:
        book.write(again, {}, target_lang="es", signature=True)
    names2, opf2, _z = opf_of(again)
    assert sum(n.endswith("jeronimo-firma.xhtml") for n in names2) == 1
    assert [r.get("idref") for r in opf2.findall(".//opf:spine/opf:itemref", OPF_NS)].count("jeronimo-firma") == 1


def test_signature_can_be_turned_off(translated):
    translated.config["firma"] = False
    out, _ = translated.export(bilingual=True)
    names, opf, _z = opf_of(out)
    assert not any("jeronimo-firma" in n for n in names)
    assert "jeronimo-firma" not in [r.get("idref") for r in opf.findall(".//opf:spine/opf:itemref", OPF_NS)]


def test_word_and_pdf_end_with_the_signature(translated, tmp_path):
    import pdfplumber
    from docx import Document

    doc = Document(str(export_docs.export_docx(translated, tmp_path / "libro.docx")))
    texts = [p.text for p in doc.paragraphs if p.text]
    assert texts[-3:] == [*credit.texts("es"), credit.URL_TEXT]
    assert len(doc.inline_shapes) == 1  # el logo
    with pdfplumber.open(export_docs.export_pdf(translated, tmp_path / "libro.pdf")) as pdf:
        last = pdf.pages[-1].extract_text() or ""
    assert "Traducido con Jerónimo" in last and credit.URL_TEXT in last

    translated.config["firma"] = False
    doc = Document(str(export_docs.export_docx(translated, tmp_path / "sin.docx")))
    assert not any("Jerónimo" in p.text for p in doc.paragraphs) and len(doc.inline_shapes) == 0
