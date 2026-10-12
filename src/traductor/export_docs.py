"""Exportar la traducción a Word (DOCX) y a PDF.

- Word: manuscrito estándar para correctores y editoriales (Times New Roman
  12, interlineado 1,5, sangría de primera línea) o bilingüe, en una tabla de
  dos columnas (original y traducción) para revisar.
- PDF: maquetado como libro (A5 o A4) con EB Garamond, títulos de capítulo y
  números de página.

Se conservan cursivas, negritas, subrayados, superíndices y llamadas a notas.
Las imágenes no se incluyen (para eso está el EPUB)."""

from __future__ import annotations

import html
from dataclasses import dataclass, field
from pathlib import Path

from . import inline
from .epub import EpubBook
from .inline import localname

ITALIC = {"i", "em", "cite", "var", "dfn"}
BOLD = {"b", "strong"}
PARAGRAPH_KINDS = {"h1": 1, "h2": 2, "h3": 3, "h4": 3, "h5": 3, "h6": 3}


@dataclass
class Run:
    text: str
    italic: bool = False
    bold: bool = False
    underline: bool = False
    sup: bool = False
    sub: bool = False
    br: bool = False


@dataclass
class Block:
    tag: str
    runs: list[Run]
    src_runs: list[Run] = field(default_factory=list)
    translated: bool = True

    @property
    def level(self) -> int:
        return PARAGRAPH_KINDS.get(self.tag, 0)


@dataclass
class Chapter:
    title: str
    blocks: list[Block]


def to_runs(marked: str, slots: dict[int, tuple[str, str]]) -> list[Run]:
    """Texto marcado → trozos con formato. ``slots`` da, para cada marca, su
    etiqueta HTML y (para los elementos fijos) su texto, p. ej. el número de nota."""
    runs: list[Run] = []
    stack: list[str] = []
    last = 0

    def add(text: str, **extra) -> None:
        if not text:
            return
        tags = set(stack)
        runs.append(Run(html.unescape(text), italic=bool(tags & ITALIC), bold=bool(tags & BOLD),
                        underline="u" in tags, sup="sup" in tags, sub="sub" in tags, **extra))

    for m, kind, n in inline._tokens(marked):
        add(marked[last:m.start()])
        last = m.end()
        tag, text = slots.get(n, ("", ""))
        if kind == "open":
            stack.append(tag)
        elif kind == "close":
            if stack:
                stack.pop()
        elif kind == "atom":
            if tag == "br":
                runs.append(Run("", br=True))
            elif text.strip():
                # Llamada a nota u otro elemento con texto: superíndice.
                runs.append(Run(text.strip(), sup=True))
    add(marked[last:])
    return runs


def collect(project, *, with_source: bool = False) -> list[Chapter]:
    """Capítulos del libro en orden de lectura, con la traducción (o el original
    donde falte) y su formato."""
    chapters: list[Chapter] = []
    with EpubBook(project.book_path) as book:
        for doc in book.documents:
            if not doc.in_spine or doc.kind != "xhtml":
                continue
            try:
                data = project.load_chapter(doc.index)
            except KeyError:
                continue
            if data.get("omitir") or not data["segmentos"]:
                continue
            _tree, elements = book._load(doc)
            blocks = []
            for seg, el in zip(data["segmentos"], elements):
                marked = inline.to_marked(el)
                slots = {k: (localname(v[1]), "".join(v[1].itertext()) if v[0] == "x" else "")
                         for k, v in marked.slots.items()}
                tgt = seg.get("tgt")
                block = Block(seg.get("tag") or localname(el), to_runs(tgt or seg["src"], slots), translated=bool(tgt))
                if with_source:
                    block.src_runs = to_runs(seg["src"], slots)
                blocks.append(block)
            chapters.append(Chapter(data.get("titulo") or "", blocks))
    return chapters


def _plain(runs: list[Run]) -> str:
    return "".join("\n" if r.br else r.text for r in runs)


# -- Word ------------------------------------------------------------------------------


def export_docx(project, out: Path, *, bilingual: bool = False, progress=None) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Pt

    chapters = collect(project, with_source=bilingual)
    cfg = project.config
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2.5))
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = None

    core = doc.core_properties
    core.title = cfg.get("titulo_traducido") or cfg.get("titulo", "")
    core.author = cfg.get("autor", "")
    core.language = project.target_lang

    def fill(paragraph, runs: list[Run]) -> None:
        for r in runs:
            if r.br:
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(r.text)
            run.italic = r.italic or None
            run.bold = r.bold or None
            run.underline = r.underline or None
            if r.sup:
                run.font.superscript = True
            if r.sub:
                run.font.subscript = True

    title = doc.add_paragraph(core.title, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if core.author:
        p = doc.add_paragraph(core.author)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for k, ch in enumerate(chapters):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if bilingual:
            doc.add_paragraph(ch.title, style="Heading 1")
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            head = table.rows[0].cells
            head[0].text, head[1].text = "Original", "Traducción"
            for cell in head:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for block in ch.blocks:
                cells = table.add_row().cells
                for cell, runs in ((cells[0], block.src_runs), (cells[1], block.runs if block.translated else [])):
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.line_spacing = 1.15
                    fill(paragraph, runs)
                    if block.level:
                        for run in paragraph.runs:
                            run.bold = True
        else:
            for block in ch.blocks:
                if block.level:
                    paragraph = doc.add_paragraph(style=f"Heading {block.level}")
                else:
                    paragraph = doc.add_paragraph()
                    fmt = paragraph.paragraph_format
                    if block.tag == "blockquote":
                        fmt.left_indent = Cm(1)
                    elif block.tag == "li":
                        fmt.left_indent = Cm(0.75)
                        paragraph.add_run("• ")
                    else:
                        fmt.first_line_indent = Cm(0.75)
                fill(paragraph, block.runs)
        if progress:
            progress((k + 1) / max(1, len(chapters)))
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


# -- PDF -------------------------------------------------------------------------------


PAGE_SIZES = {"a5": (148, 210, 10.5), "a4": (210, 297, 12)}  # ancho, alto (mm), cuerpo (pt)


def _markup(runs: list[Run]) -> str:
    """Trozos con formato → marcado de párrafo de ReportLab."""
    out = []
    for r in runs:
        if r.br:
            out.append("<br/>")
            continue
        text = html.escape(r.text, quote=False)
        if r.sup:
            text = f"<super>{text}</super>"
        if r.sub:
            text = f"<sub>{text}</sub>"
        if r.underline:
            text = f"<u>{text}</u>"
        if r.italic:
            text = f"<i>{text}</i>"
        if r.bold:
            text = f"<b>{text}</b>"
        out.append(text)
    return "".join(out)


def export_pdf(project, out: Path, *, page: str = "a5", progress=None) -> Path:
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, Frame, NextPageTemplate, PageBreak, PageTemplate, Paragraph,
                                    Spacer)

    from .cover import FONTS, unsupported_chars

    chapters = collect(project)
    missing = unsupported_chars("".join({c for ch in chapters for b in ch.blocks for c in _plain(b.runs)}))
    if missing:
        sample = " ".join(sorted(missing)[:12])
        raise ValueError(f"El PDF usa la tipografía EB Garamond, que no tiene estos caracteres: {sample}. "
                         "Exporta a Word o EPUB para este idioma.")

    for name, file in (("EBG", "EBGaramond-Regular.ttf"), ("EBG-B", "EBGaramond-SemiBold.ttf"),
                       ("EBG-I", "EBGaramond-Italic.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))
    pdfmetrics.registerFontFamily("EBG", normal="EBG", bold="EBG-B", italic="EBG-I", boldItalic="EBG-B")

    width, height, size = PAGE_SIZES.get(page, PAGE_SIZES["a5"])
    pagesize = (width * mm, height * mm)
    margin_x, margin_top, margin_bottom = (16 if page == "a5" else 25) * mm, 18 * mm, 20 * mm
    body = ParagraphStyle("body", fontName="EBG", fontSize=size, leading=size * 1.38, alignment=TA_JUSTIFY,
                          firstLineIndent=size * 1.3)
    first = ParagraphStyle("first", parent=body, firstLineIndent=0)
    quote = ParagraphStyle("quote", parent=body, leftIndent=size * 2, rightIndent=size * 1, firstLineIndent=0,
                           fontSize=size * 0.95)
    item = ParagraphStyle("item", parent=body, leftIndent=size * 1.6, firstLineIndent=-size * 0.9)
    heading = ParagraphStyle("heading", fontName="EBG-B", fontSize=size * 1.7, leading=size * 2.1,
                             alignment=TA_CENTER, spaceBefore=height * mm * 0.12, spaceAfter=size * 2.2)
    subheading = ParagraphStyle("subheading", parent=heading, fontSize=size * 1.25, leading=size * 1.6,
                                spaceBefore=size * 1.2, spaceAfter=size * 0.9)
    title_style = ParagraphStyle("title", fontName="EBG-B", fontSize=size * 2.6, leading=size * 3.1,
                                 alignment=TA_CENTER)
    author_style = ParagraphStyle("author", fontName="EBG", fontSize=size * 1.3, leading=size * 1.8,
                                  alignment=TA_CENTER, spaceBefore=size * 2)

    def numbers(canvas, doc):
        canvas.saveState()
        canvas.setFont("EBG", size * 0.85)
        canvas.drawCentredString(pagesize[0] / 2, margin_bottom / 2, str(doc.page))
        canvas.restoreState()

    frame = Frame(margin_x, margin_bottom, pagesize[0] - 2 * margin_x, pagesize[1] - margin_top - margin_bottom,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    cfg = project.config
    doc = BaseDocTemplate(str(out), pagesize=pagesize, title=cfg.get("titulo_traducido") or cfg.get("titulo", ""),
                          author=cfg.get("autor", ""), lang=project.target_lang,
                          pageTemplates=[PageTemplate("portadilla", [frame]),
                                         PageTemplate("texto", [frame], onPage=numbers)])
    story = [Spacer(1, pagesize[1] * 0.28), Paragraph(html.escape(doc.title or ""), title_style)]
    if cfg.get("autor"):
        story.append(Paragraph(html.escape(cfg["autor"]), author_style))
    story.append(NextPageTemplate("texto"))
    for k, ch in enumerate(chapters):
        story.append(PageBreak())
        after_heading = True
        for block in ch.blocks:
            text = _markup(block.runs)
            if not text.strip():
                continue
            if block.level:
                story.append(Paragraph(text, heading if block.level == 1 and after_heading else subheading))
                after_heading = True
                continue
            style = quote if block.tag == "blockquote" else item if block.tag == "li" else first if after_heading else body
            story.append(Paragraph(("• " if block.tag == "li" else "") + text, style))
            after_heading = False
        if progress:
            progress((k + 1) / max(1, len(chapters)) * 0.8)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)
    if progress:
        progress(1.0)
    return out
